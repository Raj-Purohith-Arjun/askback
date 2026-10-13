"""Excel report writer. Optional: needs `pip install openpyxl`.

Three sheets, because that is how these reports actually get used:
  Summary  the numbers someone pastes into a status update
  Cases    one row per case, filterable
  Checks   one row per check, for drilling into a failure
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from ..core.errors import AskbackError
from ..core.logging import get_logger
from ..evaluation.results import RunReport
from .redaction import DEFAULT_MODE, redacted

log = get_logger(__name__)

_CASE_HEADERS = [
    "case_id", "suite", "passed", "score", "route", "status", "risk",
    "latency_ms", "message", "reply", "tags", "error",
]
_CHECK_HEADERS = [
    "case_id", "suite", "evaluator", "check", "passed", "score", "severity", "detail",
]


def write_excel(report: RunReport, path: str | Path, *, redaction: str = DEFAULT_MODE) -> Path:
    """Write a three-sheet workbook, redacted to `redaction`."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise AskbackError("Excel reports need openpyxl. Run: pip install openpyxl") from exc

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    report = redacted(report, redaction)

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2C5FD4")
    fail_fill = PatternFill("solid", fgColor="FDE0E0")

    def style_header(sheet: Any, headers: list[str]) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    # --- Summary ----------------------------------------------------------- #
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_sheet["A1"].font = header_font
    summary_sheet["A1"].fill = header_fill
    summary_sheet["B1"].font = header_font
    summary_sheet["B1"].fill = header_fill

    for key, value in report.summary().items():
        summary_sheet.append([key, value])
    summary_sheet.append([])
    summary_sheet.append(["Suite", "Total", "Passed", "Pass rate", "Mean score"])
    for suite, data in report.by_suite().items():
        summary_sheet.append(
            [suite, data["total"], data["passed"], data["pass_rate"], data["mean_score"]]
        )
    summary_sheet.append([])
    summary_sheet.append(["Run context", ""])
    summary_sheet.append(["redaction", report.redaction])
    for key, value in report.context.items():
        summary_sheet.append([key, str(value)])

    # --- Cases -------------------------------------------------------------- #
    cases_sheet = workbook.create_sheet("Cases")
    style_header(cases_sheet, _CASE_HEADERS)
    for case in report.cases:
        cases_sheet.append([
            case.case_id, case.suite, case.passed, round(case.score, 4), case.route,
            case.status, case.risk, case.latency_ms, case.message, case.reply,
            ", ".join(case.tags), case.error or "",
        ])
        if not case.passed:
            for cell in cases_sheet[cases_sheet.max_row]:
                cell.fill = fail_fill

    # --- Checks ------------------------------------------------------------- #
    checks_sheet = workbook.create_sheet("Checks")
    style_header(checks_sheet, _CHECK_HEADERS)
    for case in report.cases:
        for result in case.results:
            for check in result.checks:
                checks_sheet.append([
                    case.case_id, case.suite, result.evaluator, check.name,
                    check.passed, round(check.score, 4), check.severity.value, check.detail,
                ])
                if not check.passed:
                    for cell in checks_sheet[checks_sheet.max_row]:
                        cell.fill = fail_fill

    for sheet, widths in (
        (summary_sheet, [28, 60]),
        (cases_sheet, [22, 14, 9, 8, 16, 12, 8, 11, 52, 70, 18, 24]),
        (checks_sheet, [22, 14, 14, 24, 9, 8, 10, 80]),
    ):
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    workbook.save(target)
    log.info("Excel report written to %s", target)
    return target
